"""產生 X01 的凌亂 workbook：合併儲存格標題、一張 sheet 多個區塊、公式、正式表格、跨區塊圖表。

python generators/x01_workbook.py  →  tasks/X01_excel_extract/inputs/產能月報.xlsx
"""
from __future__ import annotations

import datetime as dt
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference, Series
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "tasks" / "X01_excel_extract" / "inputs" / "產能月報.xlsx"
MONTHS = [f"{m}月" for m in range(1, 7)]
PLANTS = ["A廠", "B廠", "C廠"]


def build() -> Workbook:
    rnd = random.Random(20261014)
    wb = Workbook()
    ws = wb.active
    ws.title = "產能概況"

    ws.merge_cells("A1:F1")
    ws["A1"] = "2026 上半年產能概況"
    ws["A1"].font = Font(size=16, bold=True)

    # 表一：B3 標題（合併），B4:H8 區塊，第 8 列為 SUM 公式
    ws.merge_cells("B3:F3")
    ws["B3"] = "表一 各廠月產能（千片）"
    ws["B3"].alignment = Alignment(horizontal="center")
    ws.cell(4, 2, "廠區")
    for j, m in enumerate(MONTHS):
        ws.cell(4, 3 + j, m)
    for i, plant in enumerate(PLANTS):
        ws.cell(5 + i, 2, plant)
        for j in range(6):
            ws.cell(5 + i, 3 + j, rnd.randint(80, 140))
    ws.cell(8, 2, "合計")
    for j in range(6):
        col = chr(ord("C") + j)
        ws.cell(8, 3 + j, f"=SUM({col}5:{col}7)")
    ws["B10"] = "註：C廠 3 月歲修停機一週。"

    # 表二：B12 標題（未合併），B13:H16 區塊
    ws["B12"] = "表二 稼動率（%）"
    ws.cell(13, 2, "廠區")
    for j, m in enumerate(MONTHS):
        ws.cell(13, 3 + j, m)
    for i, plant in enumerate(PLANTS):
        ws.cell(14 + i, 2, plant)
        for j in range(6):
            ws.cell(14 + i, 3 + j, round(rnd.uniform(70, 98), 1))

    cats1 = Reference(ws, min_col=3, max_col=8, min_row=4)
    bar = BarChart()
    bar.title = "各廠月產能"
    for i in range(3):
        s = Series(Reference(ws, min_col=3, max_col=8, min_row=5 + i), title_from_data=False)
        s.tx = None
        bar.series.append(s)
    _name_series(bar, ws, [5, 6, 7])
    bar.set_categories(cats1)
    ws.add_chart(bar, "J3")

    line = LineChart()
    line.title = "稼動率趨勢"
    for i in range(3):
        line.series.append(Series(Reference(ws, min_col=3, max_col=8, min_row=14 + i)))
    _name_series(line, ws, [14, 15, 16])
    line.set_categories(Reference(ws, min_col=3, max_col=8, min_row=13))
    ws.add_chart(line, "J20")

    # 陷阱：同一張圖混用兩個區塊的列
    mix = LineChart()
    mix.title = "A廠產能與稼動率"
    mix.series.append(Series(Reference(ws, min_col=3, max_col=8, min_row=5)))
    mix.series.append(Series(Reference(ws, min_col=3, max_col=8, min_row=14)))
    _name_series(mix, ws, [5, 14])
    mix.set_categories(cats1)
    ws.add_chart(mix, "J37")

    # 出貨明細：正式表格 + 客戶別合計區塊（SUMIF）
    sh = wb.create_sheet("出貨明細")
    headers = ["日期", "客戶", "產品", "數量", "金額"]
    sh.append(headers)
    customers = ["好鄰居超商", "北辰汽車", "東嶺電機", "南灣零售"]
    products = ["32吋模組", "55吋模組", "車用 12.3吋"]
    start = dt.date(2026, 1, 5)
    for k in range(40):
        q = rnd.randint(50, 500)
        sh.append([start + dt.timedelta(days=4 * k), customers[k % 4], products[rnd.randint(0, 2)],
                   q, q * rnd.choice([1200, 1850, 2600])])
    tab = Table(displayName="ShipTable", ref="A1:E41")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sh.add_table(tab)
    sh["G1"] = "客戶別出貨合計"
    sh["G2"], sh["H2"] = "客戶", "數量"
    for i, c in enumerate(customers):
        sh.cell(3 + i, 7, c)
        sh.cell(3 + i, 8, f'=SUMIF($B$2:$B$41,G{3 + i},$D$2:$D$41)')
    cbar = BarChart()
    cbar.title = "客戶別出貨量"
    cbar.series.append(Series(Reference(sh, min_col=8, min_row=3, max_row=6)))
    _name_series(cbar, sh, None, fixed="H2")
    cbar.set_categories(Reference(sh, min_col=7, min_row=3, max_row=6))
    sh.add_chart(cbar, "J2")

    note = wb.create_sheet("說明")
    note["A1"] = "資料來源說明"
    note["A2"] = "產能資料由各廠每月 5 日前回報。"
    note["A3"] = "出貨明細取自出貨系統，金額單位為新台幣元。"

    fixed = dt.datetime(2026, 10, 1, 9, 0, 0)
    wb.properties.created = fixed
    wb.properties.modified = fixed
    wb.properties.creator = "ADS"
    return wb


def _name_series(chart, ws, rows, fixed=None):
    from openpyxl.chart.series import SeriesLabel
    from openpyxl.chart.data_source import StrRef
    title = ws.title
    refs = [fixed] if fixed else [f"B{r}" for r in rows]
    for s, ref in zip(chart.series, refs):
        col, row = ref[0], ref[1:]
        s.tx = SeriesLabel(strRef=StrRef(f"'{title}'!${col}${row}"))


if __name__ == "__main__":
    OUT.parent.mkdir(parents=True, exist_ok=True)
    build().save(OUT)
    print(f"wrote {OUT}")
